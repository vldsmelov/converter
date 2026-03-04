import os
import time
import uuid
from decimal import Decimal
from io import BytesIO

import requests
from openpyxl import load_workbook


KEYCLOAK_URL = os.environ.get("KEYCLOAK_URL", "http://keycloak:8080").rstrip("/")
KEYCLOAK_REALM = os.environ.get("KEYCLOAK_REALM", "uom")
KEYCLOAK_CLIENT_ID = os.environ.get("KEYCLOAK_CLIENT_ID", "uom-cli")

NSI_URL = os.environ.get("NSI_URL", "http://nsi:8000").rstrip("/")
DOCUMENTS_URL = os.environ.get("DOCUMENTS_URL", "http://documents:8000").rstrip("/")


def _wait_until(fn, timeout_s: int = 90, interval_s: float = 1.0, err: str = "timeout"):
    deadline = time.time() + timeout_s
    last_exc = None
    while time.time() < deadline:
        try:
            v = fn()
            if v:
                return v
        except Exception as e:
            last_exc = e
        time.sleep(interval_s)
    if last_exc:
        raise RuntimeError(f"{err}: last error: {last_exc}") from last_exc
    raise RuntimeError(err)


def _token(username: str, password: str) -> str:
    def _call():
        r = requests.post(
            f"{KEYCLOAK_URL}/realms/{KEYCLOAK_REALM}/protocol/openid-connect/token",
            data={
                "grant_type": "password",
                "client_id": KEYCLOAK_CLIENT_ID,
                "username": username,
                "password": password,
            },
            timeout=10,
        )
        if r.status_code != 200:
            return None
        return r.json().get("access_token")

    return _wait_until(_call, err=f"failed to get token for {username}")


def _h(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _get_json(url: str, token: str):
    r = requests.get(url, headers=_h(token), timeout=20)
    r.raise_for_status()
    return r.json()


def _post_json(url: str, token: str, payload: dict):
    r = requests.post(url, headers={**_h(token), "Content-Type": "application/json"}, json=payload, timeout=20)
    r.raise_for_status()
    return r.json()


def _ensure_uom(token: str, code: str, name: str, category_id: int, factor_to_base: str, precision: int) -> int:
    uoms = _get_json(f"{NSI_URL}/api/v1/uoms/", token)
    for u in uoms:
        if u["code"] == code:
            return int(u["id"])
    created = _post_json(
        f"{NSI_URL}/api/v1/uoms/",
        token,
        {
            "code": code,
            "name": name,
            "category": category_id,
            "factor_to_base": factor_to_base,
            "precision": precision,
        },
    )
    return int(created["id"])


def _create_item_with_rules(token: str):
    # categories
    cats = _get_json(f"{NSI_URL}/api/v1/uom-categories/", token)
    mass_id = int([c["id"] for c in cats if c["code"] == "MASS"][0])
    count_id = int([c["id"] for c in cats if c["code"] == "COUNT"][0])

    kg_id = _ensure_uom(token, "KG", "Kilogram", mass_id, "1", 3)
    ton_id = _ensure_uom(token, "TON", "Tonne", mass_id, "1000", 3)
    pcs_id = _ensure_uom(token, "PCS", "Pieces", count_id, "1", 0)
    bag_id = _ensure_uom(token, "BAG", "Bag", count_id, "1", 0)

    sku = f"E2E-ITEM-{uuid.uuid4().hex[:8]}"
    item = _post_json(
        f"{NSI_URL}/api/v1/items/",
        token,
        {
            "sku": sku,
            "name": "E2E Test Item",
            "is_active": True,
            "policy": {
                "storage_uom": kg_id,
                "posting_uom": kg_id,
                "allow_fractional": True,
                "rounding_precision": 3,
            },
        },
    )
    item_id = int(item["id"])

    # package: 1 BAG = 25 KG (ACTIVE)
    _post_json(
        f"{NSI_URL}/api/v1/packages/",
        token,
        {
            "item": item_id,
            "status": "active",
            "package_uom": bag_id,
            "content_qty": "25",
            "content_uom": kg_id,
            "supplier_code": "",
            "barcode": "",
            "effective_from": None,
            "effective_to": None,
        },
    )

    # rule: 1 PCS = 2.5 KG (ACTIVE)
    _post_json(
        f"{NSI_URL}/api/v1/rules/",
        token,
        {
            "supersedes": None,
            "item": item_id,
            "from_category": count_id,
            "to_category": mass_id,
            "rule_type": "pcs_weight",
            "conditions": {},
            "params": {"kg_per_pc": "2.5"},
            "priority": 0,
            "status": "active",
            "effective_from": None,
            "effective_to": None,
        },
    )

    return {
        "item_id": item_id,
        "sku": sku,
        "uoms": {"KG": kg_id, "TON": ton_id, "PCS": pcs_id, "BAG": bag_id},
    }


def _wait_invoice_status(token: str, invoice_id: int, want: str, timeout_s: int = 120) -> dict:
    def _poll():
        inv = _get_json(f"{DOCUMENTS_URL}/api/v1/invoices/{invoice_id}/", token)
        st = inv["status"]
        if st == "failed":
            raise RuntimeError(f"invoice failed: {inv.get('error')}")
        return inv if st == want else None

    return _wait_until(_poll, timeout_s=timeout_s, interval_s=1.0, err=f"invoice {invoice_id} did not reach {want}")


def _dec(s: str) -> Decimal:
    return Decimal(str(s))


def test_invoice_calculate_and_generate_e2e():
    # Wait basic readiness
    _wait_until(lambda: requests.get(f"{NSI_URL}/healthz", timeout=5).status_code == 200, err="NSI not ready")
    _wait_until(lambda: requests.get(f"{DOCUMENTS_URL}/healthz", timeout=5).status_code == 200, err="Documents not ready")

    operator_token = _token("operator", "operator")
    clerk_token = _token("clerk", "clerk")

    # 1) Create item + package + rule via NSI (operator)
    ctx = _create_item_with_rules(operator_token)
    item_id = ctx["item_id"]

    # 2) Create invoice via documents (clerk - no conversion roles)
    inv = _post_json(
        f"{DOCUMENTS_URL}/api/v1/invoices/",
        clerk_token,
        {
            "number": f"INV-E2E-{uuid.uuid4().hex[:6]}",
            "supplier": "ACME",
            "doc_date": "2026-03-03",
            "lines": [
                {"line_no": 1, "item_id": item_id, "qty": "1.5", "uom_code": "TON", "context": {}},
                {"line_no": 2, "item_id": item_id, "qty": "2", "uom_code": "BAG", "context": {}},
                {"line_no": 3, "item_id": item_id, "qty": "10", "uom_code": "PCS", "context": {}},
            ],
        },
    )
    invoice_id = int(inv["id"])

    # 3) Calculate
    r = requests.post(f"{DOCUMENTS_URL}/api/v1/invoices/{invoice_id}/calculate/", headers=_h(clerk_token), timeout=20)
    r.raise_for_status()

    inv_calc = _wait_invoice_status(clerk_token, invoice_id, "calculated", timeout_s=180)

    # Assert converted values line-by-line
    lines = {int(l["line_no"]): l for l in inv_calc["lines"]}

    l1 = lines[1]["converted"]
    assert _dec(l1["posting_qty"]) == Decimal("1500.000000")
    assert l1["posting_uom_code"] == "KG"

    l2 = lines[2]["converted"]
    assert _dec(l2["posting_qty"]) == Decimal("50.000000")
    assert l2["posting_uom_code"] == "KG"

    l3 = lines[3]["converted"]
    assert _dec(l3["posting_qty"]) == Decimal("25.000000")
    assert l3["posting_uom_code"] == "KG"

    # 4) Generate outputs
    r = requests.post(f"{DOCUMENTS_URL}/api/v1/invoices/{invoice_id}/generate/", headers=_h(clerk_token), timeout=20)
    r.raise_for_status()

    inv_gen = _wait_invoice_status(clerk_token, invoice_id, "generated", timeout_s=180)

    files = inv_gen.get("files") or []
    assert {f["file_type"] for f in files} == {"xlsx", "pdf"}

    xlsx_url = [f["download_url"] for f in files if f["file_type"] == "xlsx"][0]
    pdf_url = [f["download_url"] for f in files if f["file_type"] == "pdf"][0]

    # Download XLSX and validate key cells
    rx = requests.get(xlsx_url, headers=_h(clerk_token), timeout=60)
    rx.raise_for_status()
    assert "spreadsheetml" in (rx.headers.get("Content-Type") or "")

    wb = load_workbook(BytesIO(rx.content))
    ws = wb.active

    # Header row should be present
    assert ws["A1"].value == "Invoice Number"

    # Find data rows by "Line" in column A starting from row 6
    rows = []
    for r_i in range(6, 20):
        line_no = ws[f"A{r_i}"].value
        if line_no in (1, 2, 3):
            rows.append((line_no, r_i))
    assert {ln for ln, _ in rows} == {1, 2, 3}

    # Posting Qty is column E in our renderer
    r1 = dict(rows)[1]
    r2 = dict(rows)[2]
    r3 = dict(rows)[3]
    assert Decimal(str(ws[f"E{r1}"].value)) == Decimal("1500.000000")
    assert Decimal(str(ws[f"E{r2}"].value)) == Decimal("50.000000")
    assert Decimal(str(ws[f"E{r3}"].value)) == Decimal("25.000000")

    # Download PDF and validate it looks like a PDF
    rp = requests.get(pdf_url, headers=_h(clerk_token), timeout=60)
    rp.raise_for_status()
    assert (rp.headers.get("Content-Type") or "").startswith("application/pdf")
    assert rp.content[:4] == b"%PDF"
