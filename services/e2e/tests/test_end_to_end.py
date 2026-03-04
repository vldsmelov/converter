import os
import time
import uuid
from datetime import date

import httpx
from openpyxl import load_workbook


KEYCLOAK_URL = os.environ.get("KEYCLOAK_URL", "http://localhost:8080").rstrip("/")
NSI_URL = os.environ.get("NSI_URL", "http://localhost:8001").rstrip("/")
DOCS_URL = os.environ.get("DOCS_URL", "http://localhost:8002").rstrip("/")
ORIGIN = os.environ.get("ORIGIN", "http://localhost:5173")


def wait_ok(url: str, timeout_s: int = 60) -> None:
    deadline = time.time() + timeout_s
    last = None
    while time.time() < deadline:
        try:
            r = httpx.get(url, timeout=5)
            if r.status_code == 200:
                return
            last = f"{r.status_code} {r.text[:200]}"
        except Exception as e:
            last = str(e)
        time.sleep(1)
    raise RuntimeError(f"Service not ready: {url}. Last={last}")


def get_token(username: str, password: str) -> str:
    r = httpx.post(
        f"{KEYCLOAK_URL}/realms/uom/protocol/openid-connect/token",
        data={
            "grant_type": "password",
            "client_id": "uom-cli",
            "username": username,
            "password": password,
        },
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        timeout=15,
    )
    r.raise_for_status()
    return r.json()["access_token"]


def auth_headers(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def ensure_uom(client: httpx.Client, token: str, code: str, name: str, category_id: int, factor: str, precision: int):
    r = client.get(f"{NSI_URL}/api/v1/uoms/", headers=auth_headers(token))
    r.raise_for_status()
    for u in r.json():
        if u["code"] == code:
            return u["id"]

    r = client.post(
        f"{NSI_URL}/api/v1/uoms/",
        headers={**auth_headers(token), "Content-Type": "application/json"},
        json={
            "code": code,
            "name": name,
            "category": category_id,
            "factor_to_base": factor,
            "precision": precision,
        },
    )
    r.raise_for_status()
    return r.json()["id"]


def test_openapi_and_full_flow():
    # Wait for services
    wait_ok(f"{NSI_URL}/healthz")
    wait_ok(f"{DOCS_URL}/healthz")

    token = get_token("operator", "operator")

    # --- CORS + OpenAPI checks (NSI + Documents) ---
    for base in (NSI_URL, DOCS_URL):
        r = httpx.get(
            f"{base}/api/schema/",
            headers={**auth_headers(token), "Origin": ORIGIN},
            timeout=10,
        )
        assert r.status_code == 200
        assert r.headers.get("access-control-allow-origin") == ORIGIN
        data = r.json()
        assert "openapi" in data
        assert "info" in data and "title" in data["info"]

    with httpx.Client(timeout=20) as client:
        # --- Prepare reference data in NSI ---
        cats = client.get(f"{NSI_URL}/api/v1/uom-categories/", headers=auth_headers(token)).json()
        mass_id = [c["id"] for c in cats if c["code"] == "MASS"][0]
        count_id = [c["id"] for c in cats if c["code"] == "COUNT"][0]

        kg_id = ensure_uom(client, token, "KG", "Kilogram", mass_id, "1", 3)
        pcs_id = ensure_uom(client, token, "PCS", "Pieces", count_id, "1", 0)
        ton_id = ensure_uom(client, token, "TON", "Tonne", mass_id, "1000", 3)
        bag_id = ensure_uom(client, token, "BAG", "Bag", count_id, "1", 0)

        sku = f"E2E-{uuid.uuid4().hex[:8]}"
        item = client.post(
            f"{NSI_URL}/api/v1/items/",
            headers={**auth_headers(token), "Content-Type": "application/json"},
            json={
                "sku": sku,
                "name": "E2E item",
                "is_active": True,
                "policy": {
                    "storage_uom": kg_id,
                    "posting_uom": kg_id,
                    "allow_fractional": True,
                    "rounding_precision": 3,
                },
            },
        )
        item.raise_for_status()
        item_id = item.json()["id"]

        # PackageSpec: 1 BAG = 25 KG
        pkg = client.post(
            f"{NSI_URL}/api/v1/packages/",
            headers={**auth_headers(token), "Content-Type": "application/json"},
            json={
                "item": item_id,
                "package_uom": bag_id,
                "content_uom": kg_id,
                "content_qty": "25",
                "status": "active",
            },
        )
        pkg.raise_for_status()

        # Rule: 1 PCS = 2.5 KG
        rule = client.post(
            f"{NSI_URL}/api/v1/rules/",
            headers={**auth_headers(token), "Content-Type": "application/json"},
            json={
                "item": item_id,
                "from_category": count_id,
                "to_category": mass_id,
                "rule_type": "pcs_weight",
                "conditions": {},
                "params": {"kg_per_pc": "2.5"},
                "priority": 0,
                "status": "active",
            },
        )
        rule.raise_for_status()

        # --- Create invoice in Documents ---
        inv_no = f"INV-E2E-{uuid.uuid4().hex[:6]}"
        inv = client.post(
            f"{DOCS_URL}/api/v1/invoices/",
            headers={**auth_headers(token), "Content-Type": "application/json"},
            json={
                "number": inv_no,
                "supplier": "ACME",
                "doc_date": date.today().isoformat(),
                "lines": [
                    {"line_no": 1, "item_id": item_id, "qty": "1.5", "uom_code": "TON", "context": {}},
                    {"line_no": 2, "item_id": item_id, "qty": "2", "uom_code": "BAG", "context": {}},
                    {"line_no": 3, "item_id": item_id, "qty": "10", "uom_code": "PCS", "context": {}},
                ],
            },
        )
        inv.raise_for_status()
        inv_id = inv.json()["id"]

        # --- Calculate ---
        r = client.post(f"{DOCS_URL}/api/v1/invoices/{inv_id}/calculate/", headers=auth_headers(token))
        r.raise_for_status()

        # wait calculated
        for _ in range(80):
            g = client.get(f"{DOCS_URL}/api/v1/invoices/{inv_id}/", headers=auth_headers(token))
            g.raise_for_status()
            st = g.json()["status"]
            if st == "calculated":
                break
            if st == "failed":
                raise AssertionError(f"calculate failed: {g.json().get('error')}")
            time.sleep(1)
        else:
            raise AssertionError("timeout waiting for calculated")

        j = g.json()
        # verify conversions
        by_line = {ln["line_no"]: ln for ln in j["lines"]}
        assert by_line[1]["converted"]["posting_uom_code"] == "KG"
        assert float(by_line[1]["converted"]["posting_qty"]) == 1500.0
        assert float(by_line[2]["converted"]["posting_qty"]) == 50.0
        assert float(by_line[3]["converted"]["posting_qty"]) == 25.0

        # --- Generate outputs ---
        r = client.post(f"{DOCS_URL}/api/v1/invoices/{inv_id}/generate/", headers=auth_headers(token))
        r.raise_for_status()

        for _ in range(80):
            g = client.get(f"{DOCS_URL}/api/v1/invoices/{inv_id}/", headers=auth_headers(token))
            g.raise_for_status()
            st = g.json()["status"]
            if st == "generated":
                break
            if st == "failed":
                raise AssertionError(f"generate failed: {g.json().get('error')}")
            time.sleep(1)
        else:
            raise AssertionError("timeout waiting for generated")

        files = g.json()["files"]
        assert len(files) >= 2

        xlsx = [f for f in files if f["file_type"] == "xlsx"][0]
        pdf = [f for f in files if f["file_type"] == "pdf"][0]
        assert xlsx["presigned_url"], "presigned_url for xlsx should not be empty"
        assert pdf["presigned_url"], "presigned_url for pdf should not be empty"

        # --- Download via presigned URL (no auth header needed) ---
        x = client.get(xlsx["presigned_url"])
        x.raise_for_status()
        # XLSX is a zip file => starts with PK
        assert x.content[:2] == b"PK"

        wb = load_workbook(filename=bytes(x.content))
        ws = wb.active
        assert ws["A1"].value == "Invoice Number"
        assert ws["B1"].value == inv_no

        p = client.get(pdf["presigned_url"])
        p.raise_for_status()
        assert p.content[:4] == b"%PDF"
