from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Mapping

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from fpdf import FPDF

from .models import Invoice


UNICODE_FONT_CANDIDATES = [
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
    Path("/usr/local/share/fonts/DejaVuSans.ttf"),
]


def _configure_pdf_font(pdf: FPDF) -> tuple[str, bool]:
    """Return (font_name, needs_latin1_fallback)."""
    for font_path in UNICODE_FONT_CANDIDATES:
        if font_path.exists():
            pdf.add_font("DejaVuSans", "", str(font_path))
            return "DejaVuSans", False
    # Core fonts in fpdf are Latin-1 only.
    return "Helvetica", True


def _safe_pdf_text(value: object, *, latin1_only: bool) -> str:
    text = str(value) if value is not None else ""
    if not latin1_only:
        return text
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _line_item_name(line, item_name_by_id: Mapping[int, str] | None) -> str:
    ctx = getattr(line, "context", None)
    if isinstance(ctx, dict):
        raw = ctx.get("item_name")
        if isinstance(raw, str) and raw.strip():
            return raw.strip()
    if item_name_by_id:
        nm = item_name_by_id.get(int(line.item_id))
        if nm:
            return nm
    return f"item_id={line.item_id}"


def _line_posting_text(line) -> str:
    conv = getattr(line, "converted", None)
    if not conv:
        return "—"
    return f"{conv.posting_qty} {conv.posting_uom_code}"


def _line_status(line, invoice_status: str) -> str:
    conv = getattr(line, "converted", None)
    if conv:
        return "ok"
    if invoice_status == "failed":
        return "failed"
    return "—"


def render_invoice_xlsx(inv: Invoice, item_name_by_id: Mapping[int, str] | None = None) -> tuple[bytes, str, str]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Накладная"

    ws.append(["Номер накладной", inv.number])
    ws.append(["Поставщик", inv.supplier])
    ws.append(["Дата документа", inv.doc_date.isoformat() if inv.doc_date else ""])
    ws.append([])
    ws.append(["#", "Товар", "Кол-во", "ЕИ (в документе)", "Оприходование", "Статус строки"])

    for line in inv.lines.all().order_by("line_no"):
        ws.append([
            line.line_no,
            _line_item_name(line, item_name_by_id),
            str(line.qty),
            line.uom_code,
            _line_posting_text(line),
            _line_status(line, inv.status),
        ])

    col_widths = {1: 6, 2: 44, 3: 14, 4: 18, 5: 22, 6: 16}
    for col, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    bio = BytesIO()
    wb.save(bio)
    data = bio.getvalue()

    filename = f"nakladnaya_{inv.id}_{inv.number}.xlsx"
    return data, filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def render_invoice_pdf(inv: Invoice, item_name_by_id: Mapping[int, str] | None = None) -> tuple[bytes, str, str]:
    pdf = FPDF()
    pdf.add_page()
    font_name, latin1_only = _configure_pdf_font(pdf)
    pdf.set_font(font_name, size=12)

    pdf.cell(0, 8, _safe_pdf_text(f"Накладная: {inv.number}", latin1_only=latin1_only), ln=1)
    if inv.supplier:
        pdf.cell(0, 8, _safe_pdf_text(f"Поставщик: {inv.supplier}", latin1_only=latin1_only), ln=1)
    if inv.doc_date:
        pdf.cell(0, 8, _safe_pdf_text(f"Дата документа: {inv.doc_date.isoformat()}", latin1_only=latin1_only), ln=1)

    pdf.ln(4)
    pdf.set_font(font_name, size=9)

    headers = ["#", "Товар", "Кол-во", "ЕИ", "Оприходование", "Статус"]
    col_w = [10, 60, 20, 24, 36, 24]

    for i, h in enumerate(headers):
        pdf.cell(col_w[i], 7, _safe_pdf_text(h, latin1_only=latin1_only), border=1)
    pdf.ln()

    for line in inv.lines.all().order_by("line_no"):
        row = [
            str(line.line_no),
            _line_item_name(line, item_name_by_id),
            str(line.qty),
            line.uom_code,
            _line_posting_text(line),
            _line_status(line, inv.status),
        ]
        for i, val in enumerate(row):
            pdf.cell(col_w[i], 7, _safe_pdf_text(val, latin1_only=latin1_only), border=1)
        pdf.ln()

    out = pdf.output(dest="S")
    data = out if isinstance(out, (bytes, bytearray)) else out.encode("latin-1")

    filename = f"nakladnaya_{inv.id}_{inv.number}.pdf"
    return data, filename, "application/pdf"
